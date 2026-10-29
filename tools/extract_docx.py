#!/usr/bin/env python3
"""Extract reviewable content from a .docx (FS/TS): body text, headers/footers, SmartArt text,
embedded image list, tracked changes and comments. Usage: extract_docx.py file.docx [out.txt]
Needs: pandoc, python-docx. Images are listed only; view them separately when needed."""
import re, subprocess, sys, zipfile

def main(path, out):
    z = zipfile.ZipFile(path)
    parts = []
    body = subprocess.run(["pandoc", path, "-t", "plain", "--wrap=none", "--track-changes=all"],
                          capture_output=True, text=True, check=True).stdout
    parts.append("=== BODY ===\n" + body)
    import docx
    d = docx.Document(path)
    hf = []
    for i, s in enumerate(d.sections):
        for kind, h in (("HEADER", s.header), ("FOOTER", s.footer),
                        ("FIRST-HEADER", s.first_page_header), ("FIRST-FOOTER", s.first_page_footer)):
            t = [p.text for p in h.paragraphs if p.text.strip()]
            for tb in h.tables:
                for r in tb.rows:
                    t.append(" | ".join(dict.fromkeys(c.text.replace("\n", " / ") for c in r.cells)))
            if t:
                hf.append(f"section {i} {kind}: {t}")
    parts.append("=== HEADERS/FOOTERS ===\n" + "\n".join(hf))
    smart = []
    for n in sorted(x for x in z.namelist() if re.fullmatch(r"word/diagrams/data\d+\.xml", x)):
        smart.append(f"{n}: " + str(re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf8"))))
    parts.append("=== SMARTART DIAGRAMS (text only) ===\n" + ("\n".join(smart) or "none"))
    media = [f"{i.filename} ({i.file_size} bytes)" for i in z.infolist() if i.filename.startswith("word/media/")]
    parts.append("=== EMBEDDED IMAGES (not extracted) ===\n" + ("\n".join(media) or "none"))
    doc = z.read("word/document.xml").decode("utf8")
    parts.append("=== TRACKED CHANGES / COMMENTS ===\n"
                 f"insertions={doc.count('<w:ins ')} deletions={doc.count('<w:del ')} "
                 f"comments_part={'word/comments.xml' in z.namelist()}")
    text = "\n\n".join(parts)
    if out:
        open(out, "w", encoding="utf8").write(text)
        print(f"written {out}: {len(text)} chars")
    else:
        print(text)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
